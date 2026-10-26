"""Day 1: Hello documents generator for Project JARVIS.

Author: Document Engineer
Module: creator.hello_docs

Generates hello.pdf (ReportLab), hello.docx (python-docx), and hello.pptx (python-pptx)
each with the title 'Hello from JARVIS' and one line of text.
"""

from pathlib import Path
import sys
from typing import Tuple

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

import docx
from docx.shared import Pt, Inches, RGBColor

from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.dml.color import RGBColor as PptxRGBColor


def create_hello_pdf(file_path: Path) -> Path:
    """Create a minimal PDF with 'Hello from JARVIS' and one line of text using ReportLab.

    Args:
        file_path: Target path for the PDF file.

    Returns:
        Path to the generated PDF file.
    """
    try:
        doc = SimpleDocTemplate(
            str(file_path),
            pagesize=letter,
            leftMargin=54,
            rightMargin=54,
            topMargin=54,
            bottomMargin=54,
        )
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            name="HelloTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=28,
            leading=34,
            textColor=colors.HexColor("#1E3A8A"),
            alignment=1,  # Centered
            spaceAfter=20,
        )

        body_style = ParagraphStyle(
            name="HelloBody",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=14,
            leading=20,
            textColor=colors.HexColor("#1F2937"),
            alignment=1,
        )

        story = [
            Spacer(1, 40),
            Paragraph("Hello from JARVIS", title_style),
            Spacer(1, 20),
            Paragraph("This document was generated automatically by JARVIS Document Engine.", body_style),
        ]

        doc.build(story)
        return file_path
    except Exception as exc:
        raise RuntimeError(f"Failed to generate hello PDF at {file_path}: {exc}") from exc


def create_hello_docx(file_path: Path) -> Path:
    """Create a minimal Word DOCX with 'Hello from JARVIS' and one line of text using python-docx.

    Args:
        file_path: Target path for the DOCX file.

    Returns:
        Path to the generated DOCX file.
    """
    try:
        doc = docx.Document()

        title_p = doc.add_paragraph()
        title_run = title_p.add_run("Hello from JARVIS")
        title_run.font.name = "Arial"
        title_run.font.size = Pt(28)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        title_p.paragraph_format.space_after = Pt(20)

        body_p = doc.add_paragraph()
        body_run = body_p.add_run("This document was generated automatically by JARVIS Document Engine.")
        body_run.font.name = "Arial"
        body_run.font.size = Pt(14)
        body_run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        body_p.paragraph_format.space_after = Pt(12)

        doc.save(str(file_path))
        return file_path
    except Exception as exc:
        raise RuntimeError(f"Failed to generate hello DOCX at {file_path}: {exc}") from exc


def create_hello_pptx(file_path: Path) -> Path:
    """Create a minimal PowerPoint PPTX with 'Hello from JARVIS' and one line of text using python-pptx.

    Args:
        file_path: Target path for the PPTX file.

    Returns:
        Path to the generated PPTX file.
    """
    try:
        prs = Presentation()
        # 16:9 widescreen dimensions
        prs.slide_width = PptxInches(13.333)
        prs.slide_height = PptxInches(7.5)

        blank_slide_layout = prs.slide_layouts[6]  # Blank slide
        slide = prs.slides.add_slide(blank_slide_layout)

        # Background color fill
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = PptxRGBColor(0xF8, 0xFA, 0xFC)

        # Add title and text box
        tx_box = slide.shapes.add_textbox(PptxInches(1.5), PptxInches(2.5), PptxInches(10.333), PptxInches(2.5))
        tf = tx_box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = "Hello from JARVIS"
        p1.font.name = "Arial"
        p1.font.size = PptxPt(40)
        p1.font.bold = True
        p1.font.color.rgb = PptxRGBColor(0x1E, 0x3A, 0x8A)
        p1.space_after = PptxPt(20)

        p2 = tf.add_paragraph()
        p2.text = "This document was generated automatically by JARVIS Document Engine."
        p2.font.name = "Arial"
        p2.font.size = PptxPt(20)
        p2.font.color.rgb = PptxRGBColor(0x37, 0x41, 0x51)

        prs.save(str(file_path))
        return file_path
    except Exception as exc:
        raise RuntimeError(f"Failed to generate hello PPTX at {file_path}: {exc}") from exc


def generate_hello_docs(out_dir: str = "workspace") -> Tuple[str, str, str]:
    """Ensure out_dir exists, generate the three hello files, and print their paths.

    Args:
        out_dir: Directory where generated files will be placed.

    Returns:
        Tuple containing the paths (pdf_path, docx_path, pptx_path) as strings.
    """
    try:
        target_dir = Path(out_dir)
        target_dir.mkdir(parents=True, exist_ok=True)

        pdf_path = target_dir / "hello.pdf"
        docx_path = target_dir / "hello.docx"
        pptx_path = target_dir / "hello.pptx"

        created_pdf = create_hello_pdf(pdf_path)
        created_docx = create_hello_docx(docx_path)
        created_pptx = create_hello_pptx(pptx_path)

        path_pdf_str = str(created_pdf.resolve())
        path_docx_str = str(created_docx.resolve())
        path_pptx_str = str(created_pptx.resolve())

        print(f"[JARVIS Document Creator - Day 1]")
        print(f"Generated PDF:  {path_pdf_str}")
        print(f"Generated DOCX: {path_docx_str}")
        print(f"Generated PPTX: {path_pptx_str}")

        return path_pdf_str, path_docx_str, path_pptx_str
    except Exception as exc:
        print(f"Error in generate_hello_docs: {exc}", file=sys.stderr)
        raise


if __name__ == "__main__":
    print("Executing hello_docs.py verification test...")
    try:
        paths = generate_hello_docs("workspace")
        for p in paths:
            f = Path(p)
            assert f.exists(), f"File does not exist: {p}"
            assert f.stat().st_size > 0, f"File is empty: {p}"
        print("Verification successful: All 3 files generated and verified non-empty.")
    except Exception as error:
        print(f"Verification failed: {error}", file=sys.stderr)
        sys.exit(1)
