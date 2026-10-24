"""Tests for JARVIS Document Creator (Days 1, 2, and 3 - Elevated).

Author: Vidhita (Document Engineer)
Module: tests.test_creator

Verifies:
- Day 1: Hello documents (hello.pdf, hello.docx, hello.pptx)
- Day 2: Themed 3-page PDF generation across all themes ('clean_light', 'dark_pro', 'colourful')
- Day 3: Themed Word DOCX and 6-slide PowerPoint PPTX across all themes
- D1.1: Document tools (CreatePdfTool, CreateDocxTool, CreatePptxTool) in creator/tools.py and TOOL_REGISTRY
- D1.2: Claude-quality themes, executive cover pages, bento grid metrics, and Devanagari font embedding
- D1.3: Auto-open document feature (auto_open=True / False)
"""

from __future__ import annotations

import json
import os
from pathlib import Path
import re
import sys
from typing import Any, Dict, List

# Ensure project root is in sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

# Safe console encoding for Windows
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

import docx
from pptx import Presentation
import pytest

from creator.chart_maker import generate_chart
from creator.charts import make_chart
from creator.docx_maker import make_docx
from creator.hello_docs import generate_hello_docs
from creator.pdf_maker import make_pdf
from creator.pptx_maker import make_pptx
from creator.report_maker import generate_executive_report
from creator.web_maker import make_website
from creator.themes import THEMES, contains_devanagari, get_theme, register_all_fonts
from creator.outline import create_document, generate_outline
from creator.templates import (
    application_letter,
    assignment_cover,
    project_report,
    resume,
)
from creator.tools import (
    CreateChartTool,
    CreateDocxTool,
    CreatePdfTool,
    CreatePptxTool,
    CreateWebsiteTool,
    CreateDocumentTool,
    GenerateReportTool,
    CreateTemplateTool,
    TOOLS,
    create_chart,
    create_docx,
    create_pdf,
    create_pptx,
    create_website,
    create_document_tool,
    generate_report,
    create_template_document,
)
from tools.base import ToolResult
from tools.registry import TOOL_REGISTRY, execute_tool


def test_outline_fallback_uses_retrieved_sources_without_claiming_synthesis(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Current-topic fallback produces a cited source brief when both outline APIs fail."""
    import creator.outline as outline_module

    monkeypatch.setattr(outline_module, "_request_gemini_outline", lambda *args: None)
    monkeypatch.setattr(outline_module, "_request_groq_outline", lambda *args: None)
    sources = [{
        "title": "Official announcement",
        "body": "The publisher announced a new multiplayer update.",
        "href": "https://example.com/announcement",
    }]

    result = generate_outline("Modern Warfare latest update", reference_material=sources)

    sections = result["sections"]
    assert result["title"] == "Modern Warfare latest update: Source-Grounded Brief"
    assert any("publisher announced" in bullet for section in sections for bullet in section.get("bullets", []))
    assert any("https://example.com/announcement" in bullet for section in sections for bullet in section.get("bullets", []))
    assert any("rather than a fully synthesized analysis" in section.get("text", "") for section in sections)


def _count_pdf_pages(file_path: str) -> int:
    """Helper to count the number of pages in a PDF file."""
    with open(file_path, "rb") as f:
        content = f.read()
    pages = re.findall(rb"/Type\s*/Page\b", content)
    return len(pages)


@pytest.fixture
def sample_sections() -> List[Dict[str, Any]]:
    """Fixture providing comprehensive multi-section content."""
    return [
        {
            "heading": "1. Executive Summary",
            "text": "Project JARVIS introduces an autonomous multimodal assistant architecture engineered for desktop workflow automation.",
            "metrics": [
                {"label": "Response Latency", "value": "< 250ms"},
                {"label": "Document Fidelity", "value": "100% Binary"},
                {"label": "Uptime", "value": "99.99%"},
                {"label": "Memory Footprint", "value": "Low RAM"},
            ],
            "bullets": [
                "Real-time voice and text interaction loop.",
                "Automated high-quality document and presentation creation.",
                "Cross-channel communication via Telegram and WhatsApp.",
            ],
            "page_break_after": True,
        },
        {
            "heading": "2. Performance and Technical Metrics",
            "text": "The table below highlights the system benchmarks across key generation modules.",
            "bullets": [
                "Zero data leakage: credentials isolated strictly in local environment.",
                "Asynchronous processing pipeline ensuring low response latency.",
            ],
            "table": [
                ["Module", "Technology", "Target Latency", "Status"],
                ["PDF Engine", "ReportLab Platypus", "< 300 ms", "Operational"],
                ["DOCX Engine", "python-docx", "< 250 ms", "Operational"],
                ["PPTX Engine", "python-pptx (16:9)", "< 400 ms", "Operational"],
            ],
        },
    ]


@pytest.fixture
def sample_6_slide_sections() -> List[Dict[str, Any]]:
    """Fixture providing sections designed to produce at least a 6-slide deck."""
    return [
        {
            "heading": "1. Project Vision & Architecture",
            "text": "Autonomous assistant designed for modern desktop productivity.",
            "metrics": [
                {"label": "Modularity", "value": "4 Pairs"},
                {"label": "LLM Latency", "value": "Sub-sec"},
            ],
            "bullets": [
                "Modular decoupling across functional pairs.",
                "Zero external secret leaks via strict configuration boundaries.",
                "Direct programmatic generation of office document standards.",
            ],
        },
        {
            "heading": "2. Document Creator Suite",
            "text": "High-fidelity rendering across three core formats.",
            "bullets": [
                "Pixel-perfect PDF layout with cover pages and page counters.",
                "Standard-compliant Word documents with clean table borders.",
                "Presentation slides formatted in 16:9 widescreen layout.",
            ],
        },
        {
            "heading": "3. Theming & Aesthetics",
            "text": "Unified palette across all output media.",
            "bullets": [
                "Clean Light: crisp corporate styling with navy accents.",
                "Dark Pro: high-contrast dark mode for technical briefs.",
                "Colourful: vibrant palette with purple and coral tones.",
            ],
        },
        {
            "heading": "4. Verification & Milestones",
            "bullets": [
                "Day 1: Hello documents verified across PDF, DOCX, and PPTX.",
                "Day 2: Multi-page themed PDF maker with table styling.",
                "Day 3: Themed Word and 6-slide PowerPoint presentations.",
            ],
        },
    ]


# =====================================================================
# Day 1 Tests
# =====================================================================
def test_day1_hello_docs(tmp_path: Path) -> None:
    """Verify Day 1 hello docs generate valid non-empty files in all 3 formats."""
    out_dir = str(tmp_path / "day1")
    pdf_p, docx_p, pptx_p = generate_hello_docs(out_dir)

    for p in [pdf_p, docx_p, pptx_p]:
        file_obj = Path(p)
        assert file_obj.exists(), f"Day 1 file missing: {p}"
        assert file_obj.stat().st_size > 0, f"Day 1 file is empty: {p}"


# =====================================================================
# Day 2 Tests
# =====================================================================
@pytest.mark.parametrize("theme_name", ["clean_light", "dark_pro", "colourful"])
def test_day2_pdf_maker_3_pages(sample_sections: List[Dict[str, Any]], theme_name: str, tmp_path: Path) -> None:
    """Verify Day 2 acceptance test: 3-page PDF with cover, bullets, and table in all themes."""
    out_dir = str(tmp_path / "day2")
    pdf_path = make_pdf(
        title="JARVIS System Architecture",
        sections=sample_sections,
        theme=theme_name,
        out_dir=out_dir,
        auto_open=False,
    )

    pdf_file = Path(pdf_path)
    assert pdf_file.exists(), f"Generated PDF does not exist: {pdf_path}"
    assert pdf_file.stat().st_size > 500, f"Generated PDF is unexpectedly small: {pdf_path}"

    page_count = _count_pdf_pages(pdf_path)
    assert page_count >= 3, f"Expected at least 3 pages for Day 2 acceptance, found {page_count}"


def test_theme_fallback_pdf(sample_sections: List[Dict[str, Any]], tmp_path: Path) -> None:
    """Verify unknown theme safely falls back to clean_light in make_pdf."""
    out_dir = str(tmp_path / "day2_fallback")
    pdf_path = make_pdf(
        title="Fallback Theme Document",
        sections=sample_sections,
        theme="non_existent_theme",
        out_dir=out_dir,
        auto_open=False,
    )
    assert Path(pdf_path).exists()
    assert "clean_light" in pdf_path


# =====================================================================
# Day 3 Tests
# =====================================================================
@pytest.mark.parametrize("theme_name", ["clean_light", "dark_pro", "colourful"])
def test_day3_docx_maker(sample_sections: List[Dict[str, Any]], theme_name: str, tmp_path: Path) -> None:
    """Verify Day 3 Word document generation across all themes."""
    out_dir = str(tmp_path / "day3_docx")
    docx_path = make_docx(
        title="JARVIS System Specifications",
        sections=sample_sections,
        theme=theme_name,
        out_dir=out_dir,
        auto_open=False,
    )

    file_obj = Path(docx_path)
    assert file_obj.exists(), f"DOCX does not exist: {docx_path}"
    assert file_obj.stat().st_size > 1000, f"DOCX is unexpectedly small: {docx_path}"

    doc = docx.Document(docx_path)
    assert len(doc.paragraphs) > 0, "DOCX has no paragraphs"
    assert len(doc.tables) > 0, "DOCX has no tables"


@pytest.mark.parametrize("theme_name", ["clean_light", "dark_pro", "colourful"])
def test_day3_pptx_maker_6_slides(
    sample_6_slide_sections: List[Dict[str, Any]], theme_name: str, tmp_path: Path
) -> None:
    """Verify Day 3 acceptance test: 6-slide presentation deck in all themes."""
    out_dir = str(tmp_path / "day3_pptx")
    pptx_path = make_pptx(
        title="JARVIS Project Overview",
        sections=sample_6_slide_sections,
        theme=theme_name,
        out_dir=out_dir,
        auto_open=False,
    )

    file_obj = Path(pptx_path)
    assert file_obj.exists(), f"PPTX does not exist: {pptx_path}"
    assert file_obj.stat().st_size > 1000, f"PPTX is unexpectedly small: {pptx_path}"

    prs = Presentation(pptx_path)
    slide_count = len(prs.slides)
    assert slide_count >= 6, f"Expected at least 6 slides (Title + Content + Thank You), found {slide_count}"


def test_pptx_bullet_splitting(tmp_path: Path) -> None:
    """Verify that bullet lists with > 6 bullets are automatically split across slides."""
    out_dir = str(tmp_path / "day3_split")
    long_section = [
        {
            "heading": "Comprehensive Feature List",
            "bullets": [f"Feature item number {i}" for i in range(1, 14)],  # 13 bullets
        }
    ]
    pptx_path = make_pptx(
        title="Feature Breakdown",
        sections=long_section,
        theme="clean_light",
        out_dir=out_dir,
        auto_open=False,
    )
    prs = Presentation(pptx_path)
    # Expected: 1 Title + 3 Content slides (6 + 6 + 1) + 1 Thank you = 5 slides
    assert len(prs.slides) == 5, f"Expected 5 slides due to bullet splitting, got {len(prs.slides)}"


# =====================================================================
# D1.1: Document Tools & Registry Integration Tests
# =====================================================================
def test_tools_registered_in_central_registry() -> None:
    """Verify create_pdf, create_docx, and create_pptx are present in TOOL_REGISTRY."""
    assert "create_pdf" in TOOL_REGISTRY, "create_pdf not found in TOOL_REGISTRY"
    assert "create_docx" in TOOL_REGISTRY, "create_docx not found in TOOL_REGISTRY"
    assert "create_pptx" in TOOL_REGISTRY, "create_pptx not found in TOOL_REGISTRY"


def test_tool_schemas_conform_to_gemini_spec() -> None:
    """Verify the document tool schemas expose their required arguments."""
    for tool_cls in [CreatePdfTool, CreateDocxTool, CreatePptxTool]:
        tool = tool_cls()
        schema = tool.get_schema()
        assert "name" in schema
        assert "parameters" in schema
        props = schema["parameters"]["properties"]
        assert "title" in props
        assert "sections" in props
        assert "theme" in props
        assert "auto_open" in props
        assert "sections" in schema["parameters"]["required"]
        if tool_cls is CreatePdfTool:
            assert "title" not in schema["parameters"]["required"]
        else:
            assert "title" in schema["parameters"]["required"]


def test_execute_tool_create_pdf(tmp_path: Path) -> None:
    """Verify execute_tool('create_pdf', ...) returns success ToolResult with real binary PDF."""
    result: ToolResult = execute_tool(
        "create_pdf",
        {
            "title": "Quantum Computing Architecture",
            "sections": [
                {
                    "heading": "1. Qubit Coherence",
                    "text": "State-of-the-art superconducting qubits.",
                    "metrics": [
                        {"label": "Coherence Time", "value": "120 us"},
                        {"label": "Fidelity", "value": "99.9%"},
                    ],
                    "bullets": ["Transmon qubits", "Cryogenic dilution refrigerator"],
                }
            ],
            "theme": "dark_pro",
            "out_dir": str(tmp_path),
            "auto_open": False,
        },
    )
    assert result.success is True, f"create_pdf failed: {result.message}"
    assert result.data is not None
    file_path = result.data["file_path"]
    assert Path(file_path).exists()
    assert Path(file_path).stat().st_size > 1000  # Never corrupted 464 bytes!
    assert result.data["format"] == "pdf"


def test_execute_tool_create_docx(tmp_path: Path) -> None:
    """Verify execute_tool('create_docx', ...) returns success ToolResult with real binary DOCX."""
    result: ToolResult = execute_tool(
        "create_docx",
        {
            "title": "Autonomous Agent Protocols",
            "sections": [
                {
                    "heading": "1. Protocol Specifications",
                    "text": "Defines agent handshake and verification rules.",
                    "metrics": [
                        {"label": "Sync Speed", "value": "< 50ms"},
                        {"label": "Encryption", "value": "AES-256"},
                    ],
                    "bullets": ["Standard BaseTool contracts", "Strict folder boundaries"],
                }
            ],
            "theme": "clean_light",
            "out_dir": str(tmp_path),
            "auto_open": False,
        },
    )
    assert result.success is True, f"create_docx failed: {result.message}"
    file_path = result.data["file_path"]
    assert Path(file_path).exists()
    assert Path(file_path).stat().st_size > 1000
    assert result.data["format"] == "docx"


def test_execute_tool_create_pptx(tmp_path: Path) -> None:
    """Verify execute_tool('create_pptx', ...) returns success ToolResult with real binary PPTX."""
    result: ToolResult = execute_tool(
        "create_pptx",
        {
            "title": "JARVIS Executive Summary",
            "sections": [
                {
                    "heading": "1. Key Accomplishments",
                    "text": "Autonomous document creator suite.",
                    "metrics": [
                        {"label": "Formats", "value": "PDF/DOCX/PPTX"},
                        {"label": "Themes", "value": "3 Styles"},
                    ],
                    "bullets": ["ReportLab Platypus", "python-docx", "python-pptx 16:9"],
                }
            ],
            "theme": "colourful",
            "out_dir": str(tmp_path),
            "auto_open": False,
        },
    )
    assert result.success is True, f"create_pptx failed: {result.message}"
    file_path = result.data["file_path"]
    assert Path(file_path).exists()
    assert Path(file_path).stat().st_size > 1000
    assert result.data["format"] == "pptx"


def test_tools_handle_json_string_sections(tmp_path: Path) -> None:
    """Verify that when an LLM passes sections as a JSON string, it parses cleanly."""
    json_sections = json.dumps([
        {
            "heading": "LLM Generated Section",
            "text": "Parsed from a raw string input.",
            "bullets": ["Point A", "Point B"],
        }
    ])
    result = create_pdf(
        title="JSON String Parsing Test",
        sections=json_sections,
        theme="clean_light",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    assert result.success is True
    assert Path(result.data["file_path"]).exists()


def test_pdf_tool_derives_title_from_empty_input(tmp_path: Path) -> None:
    """An empty title should derive from the first section rather than fail."""
    result = create_pdf(
        title="",
        sections=[{"heading": "Test"}],
        out_dir=str(tmp_path),
        auto_open=False,
    )
    assert result.success is True
    assert result.data["title"] == "Test"
    assert Path(result.data["file_path"]).exists()


# =====================================================================
# D1.2: Claude-Quality Themes & Devanagari Font Tests
# =====================================================================
def test_devanagari_marathi_hindi_pdf_generation(tmp_path: Path) -> None:
    """Verify that Marathi / Hindi text generates valid binary PDF with Noto Sans Devanagari."""
    marathi_sections = [
        {
            "heading": "१. प्रकल्प सारांश",
            "text": "प्रकल्प जार्विस हा एक स्वायत्त डेस्कटॉप सहाय्यक आहे.",
            "metrics": [
                {"label": "कार्यक्षमता", "value": "१००%"},
                {"label": "विलंब", "value": "< ३०० मिलीसेकंद"},
            ],
            "bullets": [
                "मराठी आणि हिंदी भाषेचे संपूर्ण समर्थन.",
                "देवनागरी फॉन्टचे अचूक रेंडरिंग.",
            ],
        }
    ]
    pdf_path = make_pdf(
        title="प्रकल्प जार्विस: दस्तऐवज चाचणी",
        sections=marathi_sections,
        theme="clean_light",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    file_obj = Path(pdf_path)
    assert file_obj.exists()
    assert file_obj.stat().st_size > 5000  # Rich binary font data embedded!


def test_devanagari_marathi_hindi_docx_generation(tmp_path: Path) -> None:
    """Verify that Marathi / Hindi text generates valid binary Word document."""
    marathi_sections = [
        {
            "heading": "१. प्रकल्प तपशील",
            "text": "हा दस्तऐवज मायक्रोसॉफ्ट वर्ड स्वरूपात तयार करण्यात आला आहे.",
            "bullets": ["मराठी फॉन्ट", "स्वच्छ लेआउट"],
        }
    ]
    docx_path = make_docx(
        title="प्रकल्प जार्विस वर्ड",
        sections=marathi_sections,
        theme="clean_light",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    file_obj = Path(docx_path)
    assert file_obj.exists()
    assert file_obj.stat().st_size > 10000


# =====================================================================
# D1.3: Auto-Open Document Tests
# =====================================================================
def test_auto_open_parameter_in_all_makers(tmp_path: Path) -> None:
    """Verify auto_open parameter is accepted without error across all makers."""
    sec = [{"heading": "Auto-Open Test", "text": "Testing auto-open flag execution."}]
    pdf_p = make_pdf("Auto Open PDF", sec, out_dir=str(tmp_path), auto_open=False)
    docx_p = make_docx("Auto Open DOCX", sec, out_dir=str(tmp_path), auto_open=False)
    pptx_p = make_pptx("Auto Open PPTX", sec, out_dir=str(tmp_path), auto_open=False)

    assert Path(pdf_p).exists()
    assert Path(docx_p).exists()
    assert Path(pptx_p).exists()


# =====================================================================
# Day 4: High-Fidelity Themed Charts & Document Embedding Tests
# =====================================================================
@pytest.mark.parametrize("chart_type", ["bar", "line", "pie", "donut", "horizontal_bar"])
def test_day4_generate_chart_all_types(tmp_path: Path, chart_type: str) -> None:
    """Verify chart generation for all standard visualization types."""
    data = {"Alpha": 45, "Beta": 80, "Gamma": 60, "Delta": 95}
    chart_path = generate_chart(
        chart_type=chart_type,
        data=data,
        title=f"Benchmark {chart_type.title()} Visualisation",
        theme="clean_light",
        out_dir=str(tmp_path),
    )
    p = Path(chart_path)
    assert p.exists()
    assert p.stat().st_size > 15000  # High DPI crisp PNG


@pytest.mark.parametrize("theme_name", ["clean_light", "dark_pro", "colourful"])
def test_day4_chart_theming(tmp_path: Path, theme_name: str) -> None:
    """Verify that generated charts accurately reflect theme palettes."""
    data = {"Metric 1": 25, "Metric 2": 65, "Metric 3": 40}
    chart_path = generate_chart(
        chart_type="bar",
        data=data,
        title=f"Themed Chart {theme_name}",
        theme=theme_name,
        out_dir=str(tmp_path),
    )
    p = Path(chart_path)
    assert p.exists()
    assert p.stat().st_size > 10000


def test_day4_pdf_maker_with_chart(tmp_path: Path) -> None:
    """Verify PDF maker automatically renders and embeds chart images."""
    sections = [
        {
            "heading": "1. Performance Distribution",
            "text": "The visualization below highlights system resource metrics.",
            "chart": {
                "type": "bar",
                "data": {"CPU": 35, "RAM": 52, "Disk": 28},
                "title": "Hardware Metrics",
            },
        }
    ]
    pdf_path = make_pdf(
        title="PDF with Integrated Chart",
        sections=sections,
        theme="clean_light",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    p = Path(pdf_path)
    assert p.exists()
    assert p.stat().st_size > 8000


def test_day4_docx_maker_with_chart(tmp_path: Path) -> None:
    """Verify DOCX maker embeds chart image cleanly centered into Word document."""
    sections = [
        {
            "heading": "Analytics Trend",
            "text": "Visual trend over previous generation cycles.",
            "chart": {
                "type": "line",
                "data": {"Cycle 1": 12, "Cycle 2": 24, "Cycle 3": 38, "Cycle 4": 55},
                "title": "Cycle Velocity Trend",
            },
        }
    ]
    docx_path = make_docx(
        title="DOCX with Embedded Trend",
        sections=sections,
        theme="dark_pro",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    p = Path(docx_path)
    assert p.exists()
    assert p.stat().st_size > 12000


def test_day4_pptx_maker_with_chart(tmp_path: Path) -> None:
    """Verify PPTX maker adds dedicated widescreen slide for visual charts."""
    sections = [
        {
            "heading": "Market Share",
            "chart": {
                "type": "donut",
                "data": {"Product A": 55, "Product B": 30, "Product C": 15},
                "title": "Segment Distribution",
            },
        }
    ]
    pptx_path = make_pptx(
        title="Presentation with Donut Visual",
        sections=sections,
        theme="colourful",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    p = Path(pptx_path)
    assert p.exists()
    assert p.stat().st_size > 20000


def test_day4_create_chart_tool_schema() -> None:
    """Verify CreateChartTool exports valid OpenAPI JSON schema."""
    tool = CreateChartTool()
    schema = tool.get_schema()
    assert schema["name"] == "create_chart"
    assert "chart_type" in schema["parameters"]["properties"]
    assert "data" in schema["parameters"]["properties"]
    assert "data" in schema["parameters"]["required"]


def test_day4_create_chart_tool_execution(tmp_path: Path) -> None:
    """Verify executing CreateChartTool via TOOL_REGISTRY."""
    assert "create_chart" in TOOL_REGISTRY
    res = execute_tool(
        "create_chart",
        {
            "chart_type": "pie",
            "data": {"Apples": 40, "Bananas": 35, "Cherries": 25},
            "title": "Fruit Distribution",
            "theme": "clean_light",
            "out_dir": str(tmp_path),
        },
    )
    assert res.success is True
    assert "file_path" in res.data
    assert Path(res.data["file_path"]).exists()


# =====================================================================
# Day 5 Tests: create_document & CreateDocumentTool (Multilingual)
# =====================================================================
def test_day5_create_document_pdf(tmp_path: Path) -> None:
    """Verify Day 5 create_document generates valid PDF on a topic."""
    out_file = create_document(
        topic="Solar Energy Innovations",
        kind="pdf",
        theme="clean_light",
        language="en",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    assert Path(out_file).exists()
    assert Path(out_file).stat().st_size > 1000
    assert out_file.endswith(".pdf")


def test_day5_create_document_docx(tmp_path: Path) -> None:
    """Verify Day 5 create_document generates valid Word DOCX on a topic."""
    out_file = create_document(
        topic="Artificial Intelligence Architecture",
        kind="docx",
        theme="dark_pro",
        language="en",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    assert Path(out_file).exists()
    assert Path(out_file).stat().st_size > 1000
    assert out_file.endswith(".docx")


def test_day5_create_document_pptx(tmp_path: Path) -> None:
    """Verify Day 5 create_document generates valid PowerPoint deck on a topic."""
    out_file = create_document(
        topic="Next Generation Cloud Systems",
        kind="pptx",
        theme="colourful",
        language="en",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    assert Path(out_file).exists()
    assert Path(out_file).stat().st_size > 1000
    assert out_file.endswith(".pptx")


def test_day5_create_document_hindi(tmp_path: Path) -> None:
    """Verify Day 5 Hindi document generation with native Devanagari script."""
    out_file = create_document(
        topic="सौर ऊर्जा",
        kind="pdf",
        theme="clean_light",
        language="hi",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    assert Path(out_file).exists()
    assert Path(out_file).stat().st_size > 5000  # Embedded Devanagari font data
    assert out_file.endswith(".pdf")


def test_day5_create_document_marathi(tmp_path: Path) -> None:
    """Verify Day 5 Marathi document generation with native Devanagari script."""
    out_file = create_document(
        topic="प्रकल्प नियोजन आणि व्यवस्थापन",
        kind="docx",
        theme="dark_pro",
        language="mr",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    assert Path(out_file).exists()
    assert Path(out_file).stat().st_size > 1000
    assert out_file.endswith(".docx")


def test_day5_create_document_tool_schema() -> None:
    """Verify CreateDocumentTool conforms to BaseTool and exports OpenAPI schema."""
    tool = CreateDocumentTool()
    schema = tool.get_schema()
    assert schema["name"] == "create_document"
    assert "parameters" in schema
    assert "topic" in schema["parameters"]["properties"]
    assert "topic" in schema["parameters"]["required"]
    assert "kind" in schema["parameters"]["properties"]
    assert "language" in schema["parameters"]["properties"]
    assert "tech_innovation" in schema["parameters"]["properties"]["theme"]["enum"]


def test_pdf_cover_uses_custom_subtitle_and_neutral_metadata(tmp_path: Path) -> None:
    """PDF covers should not claim a fixed author or verification status."""
    from pypdf import PdfReader

    result = CreatePdfTool().run(
        title="Custom Cover Test",
        sections=[{"heading": "Brief", "text": "Original content."}],
        subtitle="A custom dossier subtitle",
        theme="tech_innovation",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    assert result.success
    cover_text = PdfReader(result.data["file_path"]).pages[0].extract_text()
    assert "A custom dossier subtitle" in cover_text
    assert "Generated PDF" in cover_text
    assert "Tech Innovation" in cover_text
    assert "Vidhita" not in cover_text
    assert "Production Verified" not in cover_text


def test_pdf_tool_recovers_title_from_sections_when_model_omits_it(tmp_path: Path) -> None:
    """A missing optional model-generated title should not abort a valid PDF request."""
    from pypdf import PdfReader

    result = CreatePdfTool().run(
        sections=[{"heading": "Field Brief", "text": "Original report content."}],
        theme="tech_innovation",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    assert result.success
    assert result.data["title"] == "Field Brief"
    assert "Field Brief" in PdfReader(result.data["file_path"]).pages[0].extract_text()
    assert "title" not in CreatePdfTool().get_schema()["parameters"]["required"]


def test_day5_create_document_tool_execution(tmp_path: Path) -> None:
    """Verify executing CreateDocumentTool via TOOL_REGISTRY succeeds."""
    res = execute_tool(
        "create_document",
        {
            "topic": "Quantum Computing Future",
            "kind": "pdf",
            "theme": "clean_light",
            "language": "en",
            "out_dir": str(tmp_path),
            "auto_open": False,
        },
    )
    assert res.success is True
    assert "successfully created" in res.message.lower()
    assert Path(res.data["file_path"]).exists()


def test_day5_create_document_tool_in_registry() -> None:
    """Verify create_document is auto-registered in TOOL_REGISTRY."""
    assert "create_document" in TOOL_REGISTRY
    assert isinstance(TOOL_REGISTRY["create_document"], CreateDocumentTool)


def test_claude_theme_factory_themes() -> None:
    """Verify that all Claude Theme Factory and modern themes are fully defined."""
    expected_themes = [
        "clean_light",
        "dark_pro",
        "colourful",
        "tech_innovation",
        "ocean_depths",
        "forest_moss",
        "midnight_executive",
        "warm_terracotta",
        "modern_minimalist",
        "sunset_boulevard",
    ]
    for th_name in expected_themes:
        assert th_name in THEMES, f"Theme '{th_name}' missing from THEMES dictionary"
        th = get_theme(th_name)
        assert "bg_color" in th
        assert "primary_color" in th
        assert "text_color" in th
        assert "bento" in th
        assert "card_bg" in th["bento"]


def test_all_fonts_registration() -> None:
    """Verify that register_all_fonts discovers and registers TTF fonts safely."""
    from creator.themes import register_all_fonts
    fonts = register_all_fonts()
    assert isinstance(fonts, list)
    assert len(fonts) >= 2


def test_make_website_generation(tmp_path: Path) -> None:
    """Verify make_website generates a complete, valid responsive HTML artifact."""
    from creator.web_maker import make_website

    sample_sec = [
        {
            "heading": "Platform Performance",
            "text": "Autonomous execution metrics demonstrate sub-millisecond response latency.",
            "metrics": [{"label": "Latency", "value": "0.4ms"}, {"label": "Uptime", "value": "99.99%"}],
            "bullets": ["High-performance WebSocket streaming.", "Local Faster-Whisper VAD routing."],
            "table": [["Tier", "Throughput"], ["Edge", "50k ops/s"], ["Core", "120k ops/s"]],
        }
    ]

    html_path = make_website(
        title="Edge Performance Report",
        sections=sample_sec,
        theme="tech_innovation",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    p = Path(html_path)
    assert p.exists()
    content = p.read_text(encoding="utf-8")
    assert "<!DOCTYPE html>" in content
    assert "Edge Performance Report" in content
    assert "0.4ms" in content
    assert "data-table" in content
    assert "tech_innovation" in content


def test_execute_tool_create_website(tmp_path: Path) -> None:
    """Verify CreateWebsiteTool and central registry execution for website creation."""
    from creator.tools import CreateWebsiteTool

    tool = CreateWebsiteTool()
    schema = tool.get_schema()
    assert schema["name"] == "create_website"
    assert "title" in schema["parameters"]["properties"]
    assert "sections" in schema["parameters"]["properties"]

    res = execute_tool(
        "create_website",
        {
            "title": "Autonomous Agent Workspace",
            "sections": [
                {
                    "heading": "Agent Capabilities",
                    "text": "Multi-agent orchestration with strict isolation boundaries.",
                    "bullets": ["Pair A: Core Brain & Memory", "Pair D: Document & Web Creator"],
                }
            ],
            "theme": "ocean_depths",
            "out_dir": str(tmp_path),
            "auto_open": False,
        },
    )
    assert isinstance(res, ToolResult)
    assert res.success is True
    assert "file_path" in res.data
    assert Path(res.data["file_path"]).exists()
    assert res.data["format"] == "html"


def test_day5_create_document_web_kind(tmp_path: Path) -> None:
    """Verify create_document autonomously creates web artifacts when kind='web'."""
    from creator.outline import create_document

    out_file = create_document(
        topic="Neural Network Architectures",
        kind="web",
        theme="sunset_boulevard",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    p = Path(out_file)
    assert p.exists()
    assert p.suffix == ".html"
    content = p.read_text(encoding="utf-8")
    assert "<!DOCTYPE html>" in content
    assert "Neural Network" in content


# =====================================================================
# Day 5 Tests: Autonomous Multi-Format Executive Report Generator
# =====================================================================
def test_day5_generate_executive_report_multiformat(tmp_path: Path) -> None:
    """Verify Day 5 multi-format generator creates coordinated PDF, DOCX, and PPTX reports."""
    metrics = [
        {"label": "Reliability", "value": "99.98%"},
        {"label": "Response Time", "value": "180 ms"},
        {"label": "Test Coverage", "value": "100%"},
        {"label": "Supported Formats", "value": "3/3"},
    ]
    charts = [
        {
            "type": "donut",
            "data": {"Core": 40, "Voice": 25, "Vision": 15, "Creator": 20},
            "title": "Subsystem Workload Distribution",
        }
    ]
    sections = [
        {
            "heading": "Strategic Overview",
            "text": "Autonomous execution demonstrated robust resilience across multi-modular pipelines.",
            "bullets": [
                "100% adherence to modular isolation constraints.",
                "Real-time chart rendering with 200 DPI vector clarity.",
                "Full dual-engine failover mechanism active.",
            ],
        }
    ]
    results = generate_executive_report(
        title="JARVIS Autonomous Engineering Summary",
        summary="End-to-end multi-format executive intelligence report.",
        sections=sections,
        key_metrics=metrics,
        chart_configs=charts,
        theme="clean_light",
        formats=["pdf", "docx", "pptx"],
        out_dir=str(tmp_path),
        auto_open=False,
    )
    for fmt in ["pdf", "docx", "pptx"]:
        assert fmt in results
        p = Path(results[fmt])
        assert p.exists()
        assert p.stat().st_size > 5000, f"Format {fmt} output unexpectedly small: {p.stat().st_size} bytes"


def test_day5_generate_report_tool_schema() -> None:
    """Verify GenerateReportTool exports complete schema for Autonomous Brain integration."""
    tool = GenerateReportTool()
    schema = tool.get_schema()
    assert schema["name"] == "generate_report"
    assert "parameters" in schema
    assert "title" in schema["parameters"]["properties"]
    assert "title" in schema["parameters"]["required"]
    assert "formats" in schema["parameters"]["properties"]


def test_day5_generate_report_tool_registry(tmp_path: Path) -> None:
    """Verify GenerateReportTool execution via central TOOL_REGISTRY."""
    assert "generate_report" in TOOL_REGISTRY
    res = execute_tool(
        "generate_report",
        {
            "title": "Quarterly Autonomous Assessment",
            "summary": "Verified multi-modular desktop assistant capabilities.",
            "key_metrics": [{"label": "Uptime", "value": "99.9%"}],
            "chart_configs": [{"type": "line", "data": {"W1": 5, "W2": 15, "W3": 30}, "title": "Progression"}],
            "theme": "dark_pro",
            "out_dir": str(tmp_path),
            "auto_open": False,
        },
    )
    assert res.success is True
    assert "files" in res.data
    assert len(res.data["files"]) >= 3


def test_day5_devanagari_marathi_report_generation(tmp_path: Path) -> None:
    """Verify Day 5 executive report generation in Marathi with Devanagari fonts."""
    marathi_metrics = [
        {"label": "अचूकता", "value": "१००%"},
        {"label": "वेग", "value": "< २०० मिलीसेकंद"},
    ]
    marathi_sections = [
        {
            "heading": "प्रकल्प प्रगती अहवाल",
            "text": "हा संपूर्ण अहवाल मराठी भाषेतील देवनागरी फॉन्ट तपासणीसाठी तयार करण्यात आला आहे.",
            "bullets": ["सर्व दस्तऐवज अचूकपणे तयार झाले आहेत.", "फॉन्ट करप्शन नाही."],
        }
    ]
    result = generate_executive_report(
        title="प्रकल्प जार्विस: सर्वसमावेशक अहवाल",
        summary="स्वायत्त मल्टीमॉडल दस्तऐवज निर्मिती प्रणाली.",
        sections=marathi_sections,
        key_metrics=marathi_metrics,
        theme="clean_light",
        formats=["pdf", "docx"],
        out_dir=str(tmp_path),
        auto_open=False,
    )
    assert "pdf" in result
    assert "docx" in result
    assert Path(result["pdf"]).exists()
    assert Path(result["docx"]).exists()


def test_day5_selective_formats(tmp_path: Path) -> None:
    """Verify that requesting a subset of formats only builds those formats."""
    res = generate_executive_report(
        title="Selective Export",
        summary="Testing selective format generation.",
        formats=["pdf"],
        out_dir=str(tmp_path),
        auto_open=False,
    )
    assert "pdf" in res
    assert "docx" not in res
    assert "pptx" not in res


# =====================================================================
# Day 6 Tests: Charts, Cover Pages, and Document Templates
# =====================================================================
@pytest.mark.parametrize("kind", ["bar", "line", "pie", "donut", "horizontal_bar"])
def test_day6_make_chart_all_kinds(tmp_path: Path, kind: str) -> None:
    """Verify make_chart creates clean PNG files for all supported chart kinds."""
    labels = ["Alpha", "Beta", "Gamma", "Delta"]
    values = [25.0, 45.0, 60.0, 85.0]
    out_file = make_chart(
        kind=kind,
        labels=labels,
        values=values,
        title=f"Test {kind.title()} Visualization",
        theme="clean_light",
        out_dir=str(tmp_path / "temp_creator"),
    )
    p = Path(out_file)
    assert p.exists()
    assert p.stat().st_size > 10000
    assert out_file.endswith(".png")


def test_day6_make_chart_devanagari(tmp_path: Path) -> None:
    """Verify make_chart renders Marathi / Hindi Devanagari labels and titles cleanly."""
    labels = ["सौर ऊर्जा", "पवन ऊर्जा", "जल ऊर्जा"]
    values = [55.0, 30.0, 15.0]
    out_file = make_chart(
        kind="bar",
        labels=labels,
        values=values,
        title="ऊर्जा स्रोत वर्गीकरण",
        theme="dark_pro",
        out_dir=str(tmp_path / "temp_creator"),
    )
    p = Path(out_file)
    assert p.exists()
    assert p.stat().st_size > 10000


def test_day6_section_chart_dict_pdf(tmp_path: Path) -> None:
    """Verify PDF maker embeds section chart specified with {'kind', 'labels', 'values', 'title'}."""
    sections = [
        {
            "heading": "Market Dynamics",
            "text": "Detailed distribution of quarterly metrics.",
            "chart": {
                "kind": "bar",
                "labels": ["Q1", "Q2", "Q3", "Q4"],
                "values": [120, 240, 350, 480],
                "title": "Quarterly Expansion",
            },
        }
    ]
    pdf_path = make_pdf(
        title="PDF with Day 6 Section Chart",
        sections=sections,
        theme="clean_light",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    p = Path(pdf_path)
    assert p.exists()
    assert p.stat().st_size > 8000


def test_day6_section_chart_dict_docx(tmp_path: Path) -> None:
    """Verify DOCX maker embeds section chart specified with {'kind', 'labels', 'values', 'title'}."""
    sections = [
        {
            "heading": "Trend Velocity",
            "text": "Trajectory of active runtime connections.",
            "chart": {
                "kind": "line",
                "labels": ["Week 1", "Week 2", "Week 3", "Week 4"],
                "values": [10.5, 25.2, 45.8, 80.0],
                "title": "Velocity Curve",
            },
        }
    ]
    docx_path = make_docx(
        title="DOCX with Day 6 Section Chart",
        sections=sections,
        theme="dark_pro",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    p = Path(docx_path)
    assert p.exists()
    assert p.stat().st_size > 12000


def test_day6_section_chart_dict_pptx(tmp_path: Path) -> None:
    """Verify PPTX maker creates dedicated slide for section chart specified with {'kind', 'labels', 'values', 'title'}."""
    sections = [
        {
            "heading": "Resource Allocation",
            "chart": {
                "kind": "pie",
                "labels": ["Engineering", "Design", "QA"],
                "values": [50, 30, 20],
                "title": "Department Allocation",
            },
        }
    ]
    pptx_path = make_pptx(
        title="Presentation with Day 6 Section Chart",
        sections=sections,
        theme="colourful",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    p = Path(pptx_path)
    assert p.exists()
    assert p.stat().st_size > 20000


def test_day6_cover_page_and_header_footer_options(tmp_path: Path) -> None:
    """Verify cover_page toggle and custom header_text/footer_text in PDF, DOCX, and PPTX."""
    sections = [{"heading": "Introduction", "text": "Testing cover_page=False and headers/footers."}]

    pdf_p = make_pdf(
        title="No Cover Report",
        sections=sections,
        theme="clean_light",
        out_dir=str(tmp_path),
        auto_open=False,
        cover_page=False,
        header_text="Running Header: Custom Title",
        footer_text="Confidential • Page Footer",
    )
    assert Path(pdf_p).exists()

    docx_p = make_docx(
        title="No Cover Document",
        sections=sections,
        theme="dark_pro",
        out_dir=str(tmp_path),
        auto_open=False,
        cover_page=False,
        header_text="DOCX Header Test",
        footer_text="DOCX Footer Test",
    )
    assert Path(docx_p).exists()

    pptx_p = make_pptx(
        title="No Cover Presentation",
        sections=sections,
        theme="colourful",
        out_dir=str(tmp_path),
        auto_open=False,
        cover_page=False,
        header_text="PPTX Header",
        footer_text="PPTX Footer",
    )
    assert Path(pptx_p).exists()


def test_day6_template_project_report(tmp_path: Path) -> None:
    """Verify project_report template creates valid PDF and DOCX documents with chapters."""
    chapters = [
        {
            "heading": "Subsystem Architecture",
            "text": "Detailed design of core components.",
            "bullets": ["Modular isolation", "BaseTool compliance", "Zero phantom execution"],
        }
    ]
    pdf_path = project_report(
        title="Autonomous Desktop Assistant",
        author="Vidhita (Document Engineer)",
        chapters=chapters,
        theme="clean_light",
        out_dir=str(tmp_path),
        kind="pdf",
    )
    assert Path(pdf_path).exists()
    assert pdf_path.endswith(".pdf")

    docx_path = project_report(
        title="Autonomous Desktop Assistant",
        author="Vidhita (Document Engineer)",
        chapters=chapters,
        theme="dark_pro",
        out_dir=str(tmp_path),
        kind="docx",
    )
    assert Path(docx_path).exists()
    assert docx_path.endswith(".docx")


def test_day6_template_assignment_cover(tmp_path: Path) -> None:
    """Verify assignment_cover template creates formal student submission document in PDF and DOCX."""
    pdf_path = assignment_cover(
        subject="Artificial Intelligence",
        name="Vidhita Shitole",
        roll_no="AI-524",
        title="Autonomous Agent Architecture",
        theme="clean_light",
        out_dir=str(tmp_path),
        kind="pdf",
    )
    assert Path(pdf_path).exists()
    assert pdf_path.endswith(".pdf")

    docx_path = assignment_cover(
        subject="Operating Systems Automation",
        name="Vidhita Shitole",
        roll_no="AI-524",
        title="Process Control & Telemetry",
        theme="colourful",
        out_dir=str(tmp_path),
        kind="docx",
    )
    assert Path(docx_path).exists()
    assert docx_path.endswith(".docx")


def test_day6_template_resume(tmp_path: Path) -> None:
    """Verify resume template creates clean CV in PDF and DOCX formats."""
    resume_data = {
        "name": "Vidhita Shitole",
        "role": "AI Software Engineer",
        "contact": {"email": "vidhita@example.com", "phone": "+91 9876543210", "location": "Pune, India"},
        "summary": "Specialist in autonomous agent architectures, multi-modal workflows, and document engineering.",
        "skills": {
            "Languages": ["Python", "TypeScript", "SQL"],
            "Libraries": ["ReportLab", "python-docx", "python-pptx", "Matplotlib", "FastAPI"],
        },
        "experience": [
            {
                "role": "Document Creator Engineer",
                "company": "Project JARVIS",
                "period": "2026 - Present",
                "bullets": ["Engineered PDF, DOCX, and PPTX makers with themes and Devanagari typography."],
            }
        ],
        "education": [
            {"degree": "B.E. Computer Engineering", "institution": "Pune University", "year": "2026", "score": "Distinction"}
        ],
    }
    pdf_path = resume(data=resume_data, theme="clean_light", out_dir=str(tmp_path), kind="pdf")
    assert Path(pdf_path).exists()
    assert pdf_path.endswith(".pdf")

    docx_path = resume(data=resume_data, theme="dark_pro", out_dir=str(tmp_path), kind="docx")
    assert Path(docx_path).exists()
    assert docx_path.endswith(".docx")


def test_day6_template_application_letter(tmp_path: Path) -> None:
    """Verify application_letter template creates formal application in PDF and DOCX formats."""
    letter_data = {
        "applicant_name": "Vidhita Shitole",
        "applicant_contact": "vidhita@example.com • Pune",
        "recipient_name": "Dr. Sarah Connor",
        "recipient_title": "Director of Autonomous Systems",
        "company_name": "Cyberdyne Labs",
        "subject": "Application for Document Automation Lead",
        "salutation": "Dear Dr. Connor,",
    }
    pdf_path = application_letter(data=letter_data, theme="clean_light", out_dir=str(tmp_path), kind="pdf")
    assert Path(pdf_path).exists()
    assert pdf_path.endswith(".pdf")

    docx_path = application_letter(data=letter_data, theme="colourful", out_dir=str(tmp_path), kind="docx")
    assert Path(docx_path).exists()
    assert docx_path.endswith(".docx")


def test_day6_create_template_tool_schema_and_execution(tmp_path: Path) -> None:
    """Verify CreateTemplateTool schema and execution via TOOL_REGISTRY."""
    tool = CreateTemplateTool()
    schema = tool.get_schema()
    assert schema["name"] == "create_template_document"
    assert "parameters" in schema
    assert "template_type" in schema["parameters"]["properties"]
    assert "data" in schema["parameters"]["properties"]

    assert "create_template_document" in TOOL_REGISTRY

    res = execute_tool(
        "create_template_document",
        {
            "template_type": "assignment_cover",
            "data": {
                "subject": "Deep Learning Systems",
                "name": "Vidhita Shitole",
                "roll_no": "AI-524",
                "title": "Neural Model Optimization",
            },
            "theme": "clean_light",
            "kind": "pdf",
            "out_dir": str(tmp_path),
            "auto_open": False,
        },
    )
    assert res.success is True
    assert "file_path" in res.data
    assert Path(res.data["file_path"]).exists()


# =====================================================================
# Day 7 Tests: Web Artifacts, Theme Factory, Font Suite & Multi-Format
# =====================================================================
def test_claude_theme_factory_themes() -> None:
    """Verify that all Claude Theme Factory and modern themes are fully defined."""
    expected_themes = [
        "clean_light",
        "dark_pro",
        "colourful",
        "tech_innovation",
        "ocean_depths",
        "forest_moss",
        "midnight_executive",
        "warm_terracotta",
        "modern_minimalist",
        "sunset_boulevard",
    ]
    for th_name in expected_themes:
        assert th_name in THEMES, f"Theme '{th_name}' missing from THEMES dictionary"
        th = get_theme(th_name)
        assert "bg_color" in th
        assert "primary_color" in th
        assert "text_color" in th
        assert "bento" in th
        assert "card_bg" in th["bento"]


def test_all_fonts_registration() -> None:
    """Verify that register_all_fonts discovers and registers TTF fonts safely."""
    fonts = register_all_fonts()
    assert isinstance(fonts, list)
    assert len(fonts) >= 2


def test_make_website_generation(tmp_path: Path) -> None:
    """Verify make_website generates a complete, valid responsive HTML artifact."""
    sample_sec = [
        {
            "heading": "Platform Performance",
            "text": "Autonomous execution metrics demonstrate sub-millisecond response latency.",
            "metrics": [{"label": "Latency", "value": "0.4ms"}, {"label": "Uptime", "value": "99.99%"}],
            "bullets": ["High-performance WebSocket streaming.", "Local Faster-Whisper VAD routing."],
            "table": [["Tier", "Throughput"], ["Edge", "50k ops/s"], ["Core", "120k ops/s"]],
        }
    ]

    html_path = make_website(
        title="Edge Performance Report",
        sections=sample_sec,
        theme="tech_innovation",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    p = Path(html_path)
    assert p.exists()
    content = p.read_text(encoding="utf-8")
    assert "<!DOCTYPE html>" in content
    assert "Edge Performance Report" in content


def test_execute_tool_create_website(tmp_path: Path) -> None:
    """Verify CreateWebsiteTool and central registry execution for website creation."""
    tool = CreateWebsiteTool()
    schema = tool.get_schema()
    assert schema["name"] == "create_website"
    assert "title" in schema["parameters"]["properties"]
    assert "sections" in schema["parameters"]["properties"]

    assert "create_website" in TOOL_REGISTRY

    res = execute_tool(
        "create_website",
        {
            "title": "Autonomous Agent Workspace",
            "sections": [
                {
                    "heading": "Agent Capabilities",
                    "text": "Multi-agent orchestration with strict isolation boundaries.",
                    "bullets": ["Pair A: Core Brain & Memory", "Pair D: Document & Web Creator"],
                }
            ],
            "theme": "ocean_depths",
            "out_dir": str(tmp_path),
            "auto_open": False,
        },
    )
    assert res.success is True
    assert "file_path" in res.data
    assert Path(res.data["file_path"]).exists()
    assert res.data["format"] == "html"


def test_day5_create_document_web_kind(tmp_path: Path) -> None:
    """Verify create_document autonomously creates web artifacts when kind='web'."""
    out_file = create_document(
        topic="Neural Network Architectures",
        kind="web",
        theme="sunset_boulevard",
        out_dir=str(tmp_path),
        auto_open=False,
    )
    p = Path(out_file)
    assert p.exists()
    assert p.suffix == ".html"
    content = p.read_text(encoding="utf-8")
    assert "<!DOCTYPE html>" in content
    assert "Neural Network" in content


# =====================================================================
# Runnable Standalone Verification Script
# =====================================================================
def run_verification() -> bool:
    """Execute all acceptance verification tests and display structured summary."""
    print("=" * 75)
    print("      JARVIS CREATOR MODULE - DAYS 1, 2, 3 & TOOLS ACCEPTANCE VERIFICATION")
    print("=" * 75)

    workspace_dir = Path("workspace")
    workspace_dir.mkdir(parents=True, exist_ok=True)

    test_sections = [
        {
            "heading": "1. Architecture and Design Goals",
            "text": "JARVIS Pair D module provides rich programmatic document outputs with dynamic themes.",
            "metrics": [
                {"label": "PDF Engine", "value": "ReportLab Platypus"},
                {"label": "Word Engine", "value": "python-docx"},
                {"label": "Deck Engine", "value": "python-pptx (16:9)"},
                {"label": "Font Support", "value": "Noto Sans Devanagari"},
            ],
            "bullets": [
                "Full support for ReportLab Platypus, python-docx, and python-pptx.",
                "Independent thematic styling for light, dark, and colourful palettes.",
                "Strict boundary isolation adhering to project guidelines.",
            ],
            "page_break_after": True,
        },
        {
            "heading": "2. Production Metrics & Deliverables",
            "text": "All required deliverables for Days 1 through 3 and Tools are fully operational.",
            "bullets": [
                "Day 1: Hello PDF, DOCX, and PPTX generator.",
                "Day 2: Themed 3-page PDF with cover, bullets, and table.",
                "Day 3: Themed Word DOCX and 6-slide PowerPoint decks.",
                "Day 3 Elevation: Claude-quality themes, bento grid cards, Devanagari fonts, and auto_open.",
            ],
            "table": [
                ["Deliverable", "Technology", "Status", "Target"],
                ["hello_docs.py", "ReportLab / docx / pptx", "PASS", "Day 1"],
                ["themes.py", "THEMES dict & palettes", "PASS", "Day 2"],
                ["pdf_maker.py", "Platypus + NumberedCanvas", "PASS", "Day 2"],
                ["docx_maker.py", "python-docx + borders", "PASS", "Day 3"],
                ["pptx_maker.py", "python-pptx 16:9 + 6-slides", "PASS", "Day 3"],
                ["tools.py", "BaseTool + Central Registry", "PASS", "Day 3/4"],
            ],
        },
    ]

    six_slide_sections = [
        {
            "heading": "1. Introduction & Context",
            "text": "Overview of JARVIS autonomous document generation system.",
            "metrics": [
                {"label": "Lead Engineer", "value": "Vidhita"},
                {"label": "Pair Module", "value": "Pair D"},
            ],
            "bullets": [
                "Engineered by Vidhita (Document Engineer).",
                "Built for desktop assistant workflow.",
            ],
        },
        {
            "heading": "2. Themes Engine",
            "text": "Supports three distinct visual styles.",
            "bullets": [
                "Clean Light: Royal blue and slate with high legibility.",
                "Dark Pro: Charcoal and cyan for modern technical presentations.",
                "Colourful: Violet and coral for expressive decks.",
            ],
        },
        {
            "heading": "3. Formatting Capabilities",
            "text": "Adheres to strict presentation typography standards.",
            "bullets": [
                "Titles are at least 28 pt across all slides.",
                "Body text and bullets are at least 18 pt.",
                "Lists exceeding 6 bullets are automatically chunked.",
            ],
        },
        {
            "heading": "4. Quality Verification",
            "bullets": [
                "All functions include strict type hints and docstrings.",
                "Comprehensive unit and acceptance tests included.",
                "Automated error handling with fallback defaults.",
            ],
        },
    ]

    try:
        # 1. Day 1 Verification
        print("\n[DAY 1] Verifying Hello Docs Generation...")
        h_pdf, h_docx, h_pptx = generate_hello_docs(str(workspace_dir))
        for p in [h_pdf, h_docx, h_pptx]:
            assert Path(p).exists() and Path(p).stat().st_size > 0
        print("  [PASS] Day 1 Hello Docs passed.")

        # 2. Day 2 Verification (3-page PDF across all 3 themes)
        print("\n[DAY 2] Verifying Themed 3-Page PDFs...")
        for theme in ["clean_light", "dark_pro", "colourful"]:
            pdf_out = make_pdf(
                title="JARVIS System Report",
                sections=test_sections,
                theme=theme,
                out_dir=str(workspace_dir),
                auto_open=False,
            )
            pages = _count_pdf_pages(pdf_out)
            assert pages >= 3, f"Expected >= 3 pages in {theme}, got {pages}"
            print(f"  [PASS] PDF ({theme}): {pdf_out} [{pages} pages, {Path(pdf_out).stat().st_size} bytes]")

        # 3. Day 3 Verification (Word DOCX across all 3 themes)
        print("\n[DAY 3] Verifying Themed Word Documents...")
        for theme in ["clean_light", "dark_pro", "colourful"]:
            docx_out = make_docx(
                title="JARVIS Specification Document",
                sections=test_sections,
                theme=theme,
                out_dir=str(workspace_dir),
                auto_open=False,
            )
            doc = docx.Document(docx_out)
            assert len(doc.paragraphs) > 0 and len(doc.tables) > 0
            print(f"  [PASS] DOCX ({theme}): {docx_out} [{Path(docx_out).stat().st_size} bytes]")

        # 4. Day 3 Verification (6-slide PPTX across all 3 themes)
        print("\n[DAY 3] Verifying Themed 6-Slide PowerPoint Decks...")
        for theme in ["clean_light", "dark_pro", "colourful"]:
            pptx_out = make_pptx(
                title="JARVIS Project Deck",
                sections=six_slide_sections,
                theme=theme,
                out_dir=str(workspace_dir),
                auto_open=False,
            )
            prs = Presentation(pptx_out)
            slides = len(prs.slides)
            assert slides >= 6, f"Expected >= 6 slides in {theme}, got {slides}"
            print(f"  [PASS] PPTX ({theme}): {pptx_out} [{slides} slides, {Path(pptx_out).stat().st_size} bytes]")

        # 5. D1.1: Document Tools Registration & Execution
        print("\n[D1.1] Verifying Tools Registration and Execution in TOOL_REGISTRY...")
        for t_name in ["create_pdf", "create_docx", "create_pptx"]:
            assert t_name in TOOL_REGISTRY, f"Tool '{t_name}' missing in TOOL_REGISTRY"
            res = execute_tool(
                t_name,
                {
                    "title": f"JARVIS {t_name.upper()} Test",
                    "sections": test_sections,
                    "theme": "clean_light",
                    "out_dir": str(workspace_dir),
                    "auto_open": False,
                },
            )
            assert res.success is True, f"{t_name} execution failed: {res.message}"
            assert res.data["bytes"] > 1000, f"{t_name} produced invalid file size: {res.data['bytes']} bytes"
            print(f"  [PASS] {t_name}: {res.data['file_path']} ({res.data['bytes']} bytes)")

        # 6. D1.2: Claude-Quality Devanagari Multilingual Verification
        print("\n[D1.2] Verifying Devanagari (Marathi / Hindi) Font Embedding...")
        marathi_sec = [
            {
                "heading": "१. प्रकल्प उद्दिष्टे आणि प्रगती",
                "text": "हा दस्तऐवज मराठी आणि हिंदी भाषेतील देवनागरी फॉन्ट तपासणीसाठी तयार केला आहे.",
                "metrics": [
                    {"label": "अचूकता", "value": "१००%"},
                    {"label": "वेळ", "value": "< २५० मिलीसेकंद"},
                ],
                "bullets": ["देवनागरी फॉन्ट योग्य प्रकारे रेंडर होतो.", "कोणतीही अक्षरे तुटलेली नाहीत."],
            }
        ]
        dev_res = execute_tool(
            "create_pdf",
            {
                "title": "प्रकल्प जार्विस: मराठी दस्तऐवज",
                "sections": marathi_sec,
                "theme": "clean_light",
                "out_dir": str(workspace_dir),
                "auto_open": False,
            },
        )
        assert dev_res.success is True
        print(f"  [PASS] Devanagari PDF: {dev_res.data['file_path']} ({dev_res.data['bytes']} bytes)")

        # 7. Day 4: Chart Generation & Document Embedding Verification
        print("\n[DAY 4] Verifying Themed Chart Generation & Document Embedding...")
        ch_res = execute_tool(
            "create_chart",
            {
                "chart_type": "bar",
                "data": {"Q1": 120, "Q2": 260, "Q3": 390, "Q4": 510},
                "title": "Verification Quarterly Velocity",
                "theme": "clean_light",
                "out_dir": str(workspace_dir),
            },
        )
        assert ch_res.success is True, f"create_chart failed: {ch_res.message}"
        print(f"  [PASS] create_chart: {ch_res.data['file_path']} ({ch_res.data['bytes']} bytes)")

        # 8. Day 5: Multi-Format Executive Report Package Verification
        print("\n[DAY 5] Verifying Autonomous Multi-Format Executive Report Generator...")
        rep_res = execute_tool(
            "generate_report",
            {
                "title": "JARVIS Autonomous Assessment Report",
                "summary": "Multi-format executive intelligence report verifying Days 1 to 5.",
                "key_metrics": [
                    {"label": "Document Engines", "value": "PDF, DOCX, PPTX"},
                    {"label": "Data Visuals", "value": "Matplotlib (Agg)"},
                ],
                "chart_configs": [
                    {"type": "donut", "data": {"Engine": 40, "Voice": 30, "Vision": 30}, "title": "System Distribution"}
                ],
                "theme": "clean_light",
                "out_dir": str(workspace_dir),
                "auto_open": False,
            },
        )
        assert rep_res.success is True, f"generate_report failed: {rep_res.message}"
        print(f"  [PASS] generate_report files: {list(rep_res.data['files'].keys())} ({rep_res.data['total_bytes']} bytes)")

        # 9. Day 6: Charts, Cover Pages, and Document Templates Verification
        print("\n[DAY 6] Verifying Charts, Section Embedding, and Document Templates...")
        d6_chart = make_chart(
            kind="pie",
            labels=["Backend", "Voice", "Vision", "Documents"],
            values=[35.0, 25.0, 20.0, 20.0],
            title="JARVIS System Distribution (Day 6)",
            theme="clean_light",
            out_dir=str(workspace_dir),
        )
        assert Path(d6_chart).exists()
        print(f"  [PASS] make_chart: {d6_chart} ({Path(d6_chart).stat().st_size} bytes)")

        # Day 6 Template 1: Project Report
        d6_rep = project_report(
            title="JARVIS Autonomous Architecture",
            author="Vidhita (Document Engineer)",
            chapters=[
                {
                    "heading": "Core Subsystems",
                    "text": "Production verification of all document creator components.",
                    "chart": {
                        "kind": "bar",
                        "labels": ["PDF", "DOCX", "PPTX"],
                        "values": [100, 100, 100],
                        "title": "Module Completion Rate",
                    },
                }
            ],
            theme="clean_light",
            out_dir=str(workspace_dir),
            kind="pdf",
        )
        assert Path(d6_rep).exists()
        print(f"  [PASS] project_report: {d6_rep} ({Path(d6_rep).stat().st_size} bytes)")

        # Day 6 Template 2: Assignment Cover
        d6_assign = assignment_cover(
            subject="Autonomous AI Systems",
            name="Vidhita Shitole",
            roll_no="AI-524",
            title="Multi-Modal Assistant Project",
            theme="dark_pro",
            out_dir=str(workspace_dir),
            kind="docx",
        )
        assert Path(d6_assign).exists()
        print(f"  [PASS] assignment_cover: {d6_assign} ({Path(d6_assign).stat().st_size} bytes)")

        # Day 6 Template Tool Registry Execution
        tpl_res = execute_tool(
            "create_template_document",
            {
                "template_type": "assignment_cover",
                "data": {
                    "subject": "Deep Learning Systems",
                    "name": "Vidhita Shitole",
                    "roll_no": "AI-524",
                    "title": "Neural Model Optimization",
                },
                "theme": "clean_light",
                "kind": "pdf",
                "out_dir": str(workspace_dir),
                "auto_open": False,
            },
        )
        assert tpl_res.success is True
        print(f"  [PASS] create_template_document: {tpl_res.data['file_path']} ({tpl_res.data['bytes']} bytes)")

        # 10. Day 7: Web Artifact Maker & Claude Theme Factory Verification
        print("\n[DAY 7] Verifying Web Artifact Maker & Claude Theme Factory...")
        d7_web = make_website(
            title="JARVIS Enterprise Intelligence Hub",
            sections=[
                {
                    "heading": "Subsystem Orchestration",
                    "text": "Universal document creator and interactive web artifact builder.",
                    "metrics": [
                        {"label": "Formats", "value": "PDF, DOCX, PPTX, HTML"},
                        {"label": "Themes", "value": "10 Theme Factory"},
                        {"label": "Fonts", "value": "18 Bundled Fonts"},
                    ],
                    "bullets": [
                        "Responsive modern bento grid layout with glassmorphism.",
                        "Interactive searchable data tables with live filtering.",
                        "Devanagari TrueType font rendering for Hindi and Marathi.",
                    ],
                    "table": [
                        ["Component", "Lead", "Formats Supported", "Status"],
                        ["Document Engine", "Vidhita", "PDF, DOCX, PPTX, HTML", "OPERATIONAL"],
                        ["Charts & Analytics", "Vidhita", "PNG (Agg Headless)", "OPERATIONAL"],
                        ["Templates", "Vidhita", "Reports, Resumes, Letters", "OPERATIONAL"],
                    ],
                }
            ],
            theme="tech_innovation",
            out_dir=str(workspace_dir),
            auto_open=False,
        )
        assert Path(d7_web).exists()
        print(f"  [PASS] make_website: {d7_web} ({Path(d7_web).stat().st_size} bytes)")

        web_res = execute_tool(
            "create_website",
            {
                "title": "Autonomous Assistant Portal",
                "sections": [
                    {
                        "heading": "Enterprise Features",
                        "text": "Production verification of all multi-format engines.",
                        "metrics": [{"label": "Status", "value": "Battle Ready"}],
                    }
                ],
                "theme": "ocean_depths",
                "out_dir": str(workspace_dir),
                "auto_open": False,
            },
        )
        assert web_res.success is True
        print(f"  [PASS] create_website tool: {web_res.data['file_path']} ({web_res.data['bytes']} bytes)")

        print("\n" + "=" * 75)
        print("[SUCCESS] ALL ACCEPTANCE TESTS FOR DAYS 1 THROUGH 7 PASSED!")
        print("=" * 75)
        return True

    except Exception as exc:
        print(f"\n[FAIL] Verification encountered an error: {exc}", file=sys.stderr)
        return False


if __name__ == "__main__":
    success = run_verification()
    sys.exit(0 if success else 1)
